import re,sys
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
html=open('pla-taiwan-report.html').read()
html=re.sub(r'<svg.*?</svg>','',html,flags=re.S)
html=re.sub(r'<style>.*?</style>','',html,flags=re.S)
doc=Document()
st=doc.styles['Normal']; st.font.name='Microsoft JhengHei'; st.font.size=Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
for h in ['Heading 1','Heading 2','Heading 3','Title']:
    s=doc.styles[h]; s.font.name='Microsoft JhengHei'; s.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); s.font.color.rgb=RGBColor(0x17,0x21,0x29)
sec=doc.sections[0]; sec.left_margin=sec.right_margin=Cm(2.2)
TAGC={'fact':RGBColor(0x2f,0x6b,0x43),'off':RGBColor(0x1d,0x55,0x96),'est':RGBColor(0x93,0x62,0x0b),'calc':RGBColor(0x5b,0x66,0x70)}
class P(HTMLParser):
    def __init__(s):
        super().__init__(); s.p=None; s.stack=[]; s.tag=None; s.sup=False; s.bold=False; s.table=None; s.row=None; s.cell=None; s.skip=0; s.listtype=[]; s.oln=[]; s.inref=False; s.href=None
    def newp(s,style=None):
        if s.cell is not None:
            s.p=s.cell.paragraphs[0] if not s.cell.paragraphs[0].text else s.cell.add_paragraph()
        else: s.p=doc.add_paragraph(style=style)
        return s.p
    def handle_starttag(s,t,a):
        a=dict(a); cls=a.get('class','') or ''
        if t in('nav',): s.skip+=1
        if s.skip: return
        if t=='title': s.skip+=1; return
        if t=='h1': s.newp('Title')
        elif t=='h2': s.p=doc.add_heading(level=1)
        elif t=='h3': s.p=doc.add_heading(level=2)
        elif t=='p': s.newp()
        elif t=='div' and 'cls' in cls: s.newp()
        elif t=='div' and 'meta' in cls: s.newp()
        elif t=='div' and 'legend' in cls: s.newp()
        elif t=='ul': s.listtype.append('ul')
        elif t=='ol': s.listtype.append('ol'); s.oln.append(0); s.inref='refs' in cls
        elif t=='li':
            if s.listtype and s.listtype[-1]=='ol':
                s.oln[-1]+=1; s.newp(); s.p.add_run('%d. '%s.oln[-1]) if not s.inref else s.p.add_run('[%d] '%s.oln[-1])
            else: s.newp('List Bullet')
        elif t=='span' and 'tag' in cls.split():
            s.tag=cls.split()[1] if len(cls.split())>1 else None
        elif t=='sup': s.sup=True
        elif t in('b','strong','th'): s.bold=True
        elif t=='a': s.href=a.get('href')
        elif t=='figure':
            fid=a.get('id'); doc.add_picture('png/%s.png'%fid, width=Cm(16.5)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
            s.skip_fig=True
        elif t=='figcaption': s.newp(); s.p.alignment=WD_ALIGN_PARAGRAPH.CENTER; s.cap=True
        elif t=='table':
            s.table=doc.add_table(rows=0,cols=0); s.table.style='Table Grid'; s.rows=[]
        elif t=='tr': s.rows.append([])
        elif t in('td','th'): s.rows[-1].append([])
    def handle_endtag(s,t):
        if t=='nav' and s.skip: s.skip-=1; return
        if t=='title' and s.skip: s.skip-=1; return
        if s.skip: return
        if t=='span': s.tag=None
        elif t=='sup': s.sup=False
        elif t in('b','strong','th'): s.bold=False
        elif t in('ul','ol'): s.listtype.pop(); 
        elif t=='ol': pass
        elif t=='a' and s.href and s.inref and s.href.startswith('http') and s.p is not None:
            r=s.p.add_run(' '+s.href); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(0x1d,0x55,0x96); s.href=None
        elif t=='table':
            rows=s.rows; n=max(len(r) for r in rows)
            tb=s.table; 
            for c in range(n): tb.add_column(Cm(4))
            for i,r in enumerate(rows):
                cells=tb.add_row().cells
                for j,txt in enumerate(r):
                    par=cells[j].paragraphs[0]; run=par.add_run(''.join(txt).strip()); run.font.size=Pt(9.5); run.bold=(i==0)
            s.table=None
        elif t=='figcaption': s.cap=False
    def handle_data(s,d):
        if s.skip: return
        if s.table is not None:
            if s.rows and s.rows[-1]: 
                if s.sup: d='['+d+']'
                if s.tag: d='【'+d+'】'
                s.rows[-1][-1].append(d)
            return
        if not d.strip() and not (s.p is not None and d==' '): return
        if s.p is None: return
        if s.tag: r=s.p.add_run('【%s】'%d.strip()); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=TAGC.get(s.tag,RGBColor(0,0,0)); return
        if s.sup: r=s.p.add_run('[%s]'%d); r.font.superscript=True; r.font.size=Pt(8); return
        if s.inref and s.href: return
        r=s.p.add_run(re.sub(r'\s+',' ',d)); r.bold=s.bold
        if getattr(s,'cap',False): r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0x56,0x63,0x6e)
P().feed(html)
pass
doc.save('解放軍對台軍事行動態勢研析.docx'); print('saved')
