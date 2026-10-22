# -*- coding: utf-8 -*-
"""由 data.py 與 figures/*.png 產生 Word 檔。用法（於報告資料夾）：python3 src/todocx.py"""
import os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from data import REFS, CHAPTERS, E, TIMELINE
from build import TITLE, DATE, LABEL, PHOTOS_META

FONT = 'Microsoft JhengHei'
COL = {'fact': RGBColor(0x2f, 0x6b, 0x43), 'off': RGBColor(0x1d, 0x55, 0x96), 'est': RGBColor(0x93, 0x62, 0x0b),
       'calc': RGBColor(0x5b, 0x66, 0x70), 'gap': RGBColor(0x7a, 0x4f, 0xa0)}

doc = Document()
st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
for h in ['Heading 1', 'Heading 2', 'Heading 3', 'Title']:
    s = doc.styles[h]; s.font.name = FONT; s.element.rPr.rFonts.set(qn('w:eastAsia'), FONT); s.font.color.rgb = RGBColor(0x17, 0x21, 0x29)
sec = doc.sections[0]; sec.left_margin = sec.right_margin = Cm(2.2)


def para(text='', style=None, size=None, color=None, bold=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text); r.bold = bold
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return p


def label(p, k):
    r = p.add_run(f'［{LABEL[k]}］'); r.font.color.rgb = COL[k]; r.bold = True


def refs(p, rs):
    if rs:
        r = p.add_run(' [' + ','.join(map(str, rs)) + ']'); r.font.size = Pt(8.5); r.font.color.rgb = COL['off']


def fig(n, cap):
    doc.add_picture(os.path.join(OUT, 'figures', f'圖{n:02d}.png'), width=Cm(16.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(cap, size=9, color=COL['calc']); p.alignment = WD_ALIGN_PARAGRAPH.CENTER


doc.add_heading(TITLE, 0)
para('無人機、機器狗與無人地面載具、無人艇與潛航器、反無人機與電子干擾體系，以及可能危害直升機飛行的武器與作為', color=COL['calc'])
para(f'資料截止：{DATE}｜條目 {len(E)} 筆｜出處 {len(REFS)} 筆｜公開資訊彙整，非機密情資', size=9, color=COL['calc'])
p = para('判讀層級：')
for k, d in (('fact', '已發生、公開可查證'), ('off', '官方或官媒之說法'), ('est', '智庫、媒體或本報告之分析'), ('gap', '公開資料未見')):
    label(p, k); p.add_run(d + '　')

doc.add_heading('總覽：分類與條目索引', 1)
fig(1, '圖 01　解放軍無人系統與直升機威脅分類總覽（括號內為條目數）')
t = doc.add_table(rows=1, cols=4); t.style = 'Table Grid'
for c, h in zip(t.rows[0].cells, ['編號', '名稱', '類型', '狀態']):
    c.paragraphs[0].add_run(h).bold = True
for e in E:
    cells = t.add_row().cells
    for c, v in zip(cells, [e['id'], e['name'], e['kind'], e['status']]):
        c.paragraphs[0].add_run(v).font.size = Pt(9)

for a, no, title, intro in CHAPTERS:
    doc.add_heading(f'{no}　{title}', 1)
    para(intro)
    if a == 'c5':
        fig(2, '圖 02　可威脅直升機之主要武器最大射程與最低攔截高度（公開資料）')
    sub = None
    for e in E:
        if e['ch'] != a:
            continue
        if e['sub'] != sub:
            sub = e['sub']; doc.add_heading(sub, 2)
        h = doc.add_heading(f"{e['id']}　{e['name']}", 3)
        p = para(); label(p, e['label'])
        meta = '｜'.join(x for x in [e['alias'], e['kind'], '研製：' + e['maker'], '狀態：' + e['status']] if x)
        r = p.add_run(meta); r.font.size = Pt(9.5)
        for k, lab in (('specs', '主要性能'), ('activity', '動態與部署')):
            if e[k]:
                p = para(); p.add_run(lab + '：').bold = True; p.add_run(e[k])
        if e['assess']:
            p = para(); p.add_run('研判意涵：').bold = True; label(p, e['alabel']); p.add_run(e['assess'])
        p = para('出處', size=8.5, color=COL['calc']); refs(p, e['refs'])
        if not e['refs']:
            p.add_run('：無').font.size = Pt(8.5)

doc.add_heading('第六章　對台相關活動時間線', 1)
fig(3, '圖 03　解放軍無人系統對台相關活動時間線（2022 至 2026 年）')
for d, ev, rs in TIMELINE:
    p = para(); p.add_run(d + '　').bold = True; p.add_run(ev); refs(p, rs)

doc.add_heading('結論：研判與觀察指標', 1)
for txt in ['解放軍無人系統已形成高空偵察、中空長航時察打、隱身僚機、蜂群母機、巡飛彈、地面與水下無人載具的完整譜系，多數仍在試驗或展示階段，但對台偵察與電子情報蒐集已常態化。',
            '開戰初期可能以殲-6 改裝無人機、反輻射巡飛彈與蜂群消耗並壓制台灣防空雷達，再由有人機與攻擊-11 類無人機跟進打擊。',
            '反無人機已升格為獨立作戰體系，雷射、高功率微波與干擾的目標是削弱台灣以無人機反登陸的效果。',
            '直升機在灘頭、近岸與城鎮低空作業時，同時面對便攜防空飛彈、彈砲合一系統、FPV 與巡飛彈，以及大範圍電子干擾；平時則可能遭遇熱焰彈、近距離逼近與雷射照射等灰色地帶干擾。']:
    p = para(style='List Bullet'); label(p, 'est'); p.add_run(txt)
doc.add_heading('觀察指標', 2)
for txt in ['BZK-005、TB-001、攻擊-2 以外的新機型（如攻擊-11、彩虹-7、九天）是否出現在台灣周邊空域。', '殲-6 改裝無人機的部署基地與數量變化。',
            '076 型艦服役與攻擊-21 上艦進度。', '機器狗、無人車在東部戰區登陸演練中的運用頻率。',
            'LY-1、颶風-3000 等反無人機系統是否部署到沿海機場、港口或兩棲艦。', '解放軍艦機對外國直升機施放熱焰彈、雷射照射或危險逼近事件是否擴及台灣周邊。']:
    para(txt, style='List Bullet')

if PHOTOS_META:
    names = {e['id']: e['name'] for e in E}
    doc.add_heading('附錄一　裝備圖片', 1)
    para(f'以下為本資料庫提及裝備中，可在 Wikimedia Commons 找到自由授權照片者，共 {len(PHOTOS_META)} 張。每張圖註明作者與授權；部分照片為航展或閱兵展示機，外觀可能與服役型號不同。')
    for ph in PHOTOS_META:
        doc.add_picture(os.path.join(OUT, 'figures', 'photos', ph['img']), width=Cm(11))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = para(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{ph['id']}　{names[ph['id']]}"); r.bold = True; r.font.size = Pt(9.5)
        if ph['caption']:
            p.add_run('：' + ph['caption']).font.size = Pt(9.5)
        p = para(f"圖片：{ph['author'] or '未署名'}，{ph['license']}，Wikimedia Commons {ph['page']}", size=8, color=COL['calc'])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('附錄二　參考資料' if PHOTOS_META else '參考資料', 1)
for k in sorted(REFS):
    t, u = REFS[k]
    p = para(f'[{k}] {t}', size=9); r = p.add_run(' ' + u); r.font.size = Pt(8); r.font.color.rgb = COL['off']
para('本資料庫僅彙整公開資訊；部分外文來源係經新聞或百科轉述，正式引用前建議查閱原始報告。「推估」不代表已知作戰計畫。', size=9, color=COL['calc'])

path = os.path.join(OUT, TITLE + '.docx')
doc.save(path); print('saved', path)
